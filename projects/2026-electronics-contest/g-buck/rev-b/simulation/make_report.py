"""Create reports and reviewable artifacts from completed simulation results."""
from pathlib import Path
import json,csv,zipfile,html
import numpy as np
from simulate import ROOT,DATA,FIGURES,REPORTS,DOWNLOADS
import matplotlib.pyplot as plt

def main():
    r=json.loads((DATA/'summary.json').read_text())
    ss=json.loads((DATA/'sensitivity_summary.json').read_text())
    by={x['case']:x for x in r+ss}
    p=by['vin25_r8']; c=by['convergence_50ns']
    budget=p['pout']/.85-p['pin_model']
    for name in ['load_steps','line_step','load_release','short_no_software']:
        assert 'fault_time' in by[name], 'Run the updated dynamic simulations first'
    # Human-reviewable waveforms, sampled traces and event detail.
    fig,ax=plt.subplots(3,1,figsize=(10,9),layout='constrained')
    for name,label in [('load_steps','Load 20 -> 8 -> 5 ohm'),('line_step','Input 15 -> 25 V'),('load_release','Load 8 ohm -> open')]:
        a=np.load(DATA/f'{name}.npz')['trace']
        for x,col in zip(ax,[1,2,4]):x.plot(a[:,0]*1000,a[:,col],label=label)
        b=a[::max(1,len(a)//3000)]
        np.savetxt(DATA/f'{name}_sampled.csv',b,delimiter=',',header='t,vout,il,vin,duty,iref,vext,isource,vsw,vcap,latch',comments='')
    for x,y in zip(ax,['Output (V)','Inductor current (A)','Duty command']):
        x.set(xlabel='Time (ms)',ylabel=y);x.grid();x.legend(fontsize=8)
    fig.savefig(FIGURES/'transients.png',dpi=160);plt.close(fig)
    a=np.load(DATA/'short_no_software.npz')['trace']
    fig,ax=plt.subplots(2,1,figsize=(10,6),layout='constrained')
    ax[0].plot(a[:,0]*1000,a[:,1]);ax[0].set_ylabel('Output (V)')
    ax[1].plot(a[:,0]*1000,a[:,2],label='IL (A)');ax[1].plot(a[:,0]*1000,a[:,10],label='Fault latch');ax[1].legend()
    for x in ax:x.set(xlim=(79.8,83),xlabel='Time (ms)');x.grid()
    fig.savefig(FIGURES/'short.png',dpi=160);plt.close(fig)
    a=np.load(DATA/'vin25_r8.npz')['last']
    fig,ax=plt.subplots(2,1,figsize=(10,5),layout='constrained')
    ax[0].plot((a[:,0]-a[0,0])*1e6,a[:,2]);ax[0].set_ylabel('Inductor current (A)')
    ax[1].plot((a[:,0]-a[0,0])*1e6,a[:,1]);ax[1].set_ylabel('Output (V)')
    for x in ax:x.set_xlabel('Time in final 100 us (us)');x.grid()
    fig.savefig(FIGURES/'ripple.png',dpi=160);plt.close(fig)
    aux=np.linspace(0,1.2,121); extra=np.linspace(0,.8,81)
    eta=100*p['pout']/(p['pin_model']+aux[None,:]+extra[:,None])
    fig,ax=plt.subplots(figsize=(8,5),layout='constrained')
    im=ax.contourf(aux,extra,eta,levels=np.arange(75,97,1));fig.colorbar(im,label='Total efficiency (%)')
    ax.contour(aux,extra,eta,levels=[85],colors='red');ax.set(xlabel='Auxiliary input power incl. gate drive (W)',ylabel='Additional switching/core/thermal losses (W)',title='25 V / 8 ohm: 85% boundary (assumed model)')
    fig.savefig(FIGURES/'efficiency_budget.png',dpi=160);plt.close(fig)
    # Dynamic ranges and timing are measured from saved waveforms, not guessed.
    dyn={}
    for name in ['load_steps','line_step','load_release','short_no_software']:
        a=np.load(DATA/f'{name}.npz')['trace'];m=a[:,0]>.08
        dyn[name]={'post_vmin':float(a[m,1].min()),'post_vmax':float(a[m,1].max()),
                   'post_il_min':float(a[m,2].min()),'post_il_max':float(a[m,2].max())}
    (DATA/'dynamic_analysis.json').write_text(json.dumps(dyn,indent=2))
    cv0=by['settled_15V_open'];cv15=by['settled_15V_20ohm'];cv25=by['settled_25V_20ohm']
    table='| 输入V | 外部负载Ω | 输出均值V | 外部电流A | 电感电流范围A | 故障锁存 |\n|---|---|---|---|---|---|\n'
    for x in r[:12]:
        table+=f"| {x['vin']} | {'空载' if x['load']>1e10 else x['load']} | {x['vout_mean']:.4f} | {x['i_load_mean']:.4f} | {x['il_min']:.3f}～{x['il_max']:.3f} | {'是' if x['fault_latched'] else '否'} |\n"
    findings=f'''# G题同步Buck Rev B 仿真验证报告

日期：2026-10-21。依据：用户提供的Rev B完整原理图PDF（10页）、2026年电子电赛试题DOCX中的G题、B3/Rev B变化标注版物料清单DOCX。源文件SHA256见references/source_files.json。

## 1. 结论及适用范围

PDF重建的开关电路在本报告拟议的连续双环控制下，可满足题目第(1)～(6)项的标称稳态电压、电流要求；这是“电路方案在明确假设下可行”的证据，不能认定现有STM32程序已通过，用户未提供程序。仿真使用ngspice 44.2，执行12种输入/负载组合、4项动态工况、步长收敛、2项频率敏感性、3项延长稳态检查，以及参考/比较极性DC扫描。

总效率≥85%尚未验证。物料清单没有给出M1/M2的具体效率曲线、核心板与屏幕耗电、电感磁芯损耗、电容ESR及厂家MOSFET模型，不能把理想化功率级效率当成整机成绩。第(8)显示及第(9)按键设置未验证。未执行原生EDA ERC、PCB DRC、自举/驱动详细仿真或热分析。

## 2. 原理图与物料清单核对

关键参数一致：IRF540NPbF两只、IR2104PbF、220uH/Isat≥3A/Irms≥1.5A/DCR≤0.15Ω、输入与输出50mΩ、电感25mΩ、470uF+10uF输出电容、INA240A1D双向采样、INA181A1DBV负载采样、LM393保护及HC74/HC08锁存使能。

BOM明确R11=23kΩ 1%、R12=10kΩ 1%；不能把过压门限当成0.1%精度。C4的直流偏压有效容量、自举电容有效容量、电感实物参数、辅助模块具体型号仍待确认。本次没有改变原理图或物料清单，也没有凭空增加47ns/250pF对应的采购器件。

## 3. 模型及控制假设

本次手工按PDF建立SPICE网表，不是从EDA直接导出。功率开关用44mΩ双向受控开关及近似体二极管表示；D1用近似肖特基模型。电感DCR取0.15Ω上限，C3 ESR假设80mΩ；保险丝/连接等效20mΩ。PWM=100kHz、死区520ns、控制边沿RC=47ns均为模型假设，不是实测或厂家保证。未包含真实MOSFET栅荷、米勒平台、Coss、反向恢复、寄生电感、饱和、温升及磁芯损耗。

拟议控制：电压环PI Kp=0.5A/V、Ki=50A/(V·s)，输出电感目标电流−0.3～1A；电感电流环PI Kp=0.15/A、Ki=200/(A·s)，加入Vcap/Vin前馈、占空比0.08～0.88和积分抗饱和。1ms低侧预充后软启动，10ms达到10V目标，斜坡期间加0.48A电容充电前馈。内环反馈含1kΩ/4.7nF滤波。

此连续控制不等价于F103C8T6的离散固件。空载采用有限负电流主动调节，回收能量暂存在输入电容；真实程序若采用跳脉冲或禁止负电流，需要重新验证。初始试算表明，简单禁止负电流且不处理斜坡积分会产生空载过冲，最终模型已加入上述处理；这属于控制方案选择，不能据此判定原图存在硬件缺陷。

硬件保护在动态模型中采用标称±1.8A/11V、0.4us采样RC和1.3us比较输出RC及故障锁存行为近似。这里的RC不等于实际传播延迟；真实INA240、LM393、逻辑门和IR2104故障链路尚需验证。上电电源斜率、POR、ARM竞争、自举VB−VS及UVLO未验证。

## 4. 稳态结果

下表为40ms主矩阵末5ms的时间加权均值；电流只计外部电阻，保留板载2.2kΩ泄放与11kΩ分压。峰峰包络包含尚未完全消退的控制误差，不能全部当成开关纹波。

{table}

关键恒压工况另延长至80ms：15V空载{cv0['vout_mean']:.6f}V；15V/20Ω={cv15['vout_mean']:.6f}V；25V/20Ω={cv25['vout_mean']:.6f}V。小数仅表示数值计算结果，不表示实物精度。

对应负载调整差={abs(cv0['vout_mean']-cv15['vout_mean']):.6f}V，源调整差={abs(cv25['vout_mean']-cv15['vout_mean']):.6f}V，均低于0.5V限值。25V/8Ω={p['i_load_mean']:.6f}A；25V/5Ω={by['vin25_r5']['i_load_mean']:.6f}A；两者差={abs(p['i_load_mean']-by['vin25_r5']['i_load_mean']):.6f}A，满足1±0.1A及差值≤0.2A的稳态要求。电感目标1A包含泄放/分压电流，故外部负载电流略低于1A。

## 5. 动态、保护及瞬时限流边界

动态测试均从80ms已稳定状态开始：输入15→25V的变化时间10us；20→8Ω负载在80ms切换，8→5Ω在110ms切换；另测试8Ω卸载至空载、8Ω突变为10mΩ短路。理想开关负载用于压力测试，不等同于题目缓慢调节滑线电阻。

20→8→5Ω全过程外部负载峰值={by['load_steps']['i_load_peak_all']:.3f}A。输出电容维持突变前电压，因此不能保证任何时刻都≤1A；这是电容直接连接输出的拓扑边界。若题目“最大不超过1A”按瞬时严格判定，当前方案不能仅凭稳态合格宣布达标，需要明确测试口径，或另做输出侧限流/电容能量约束设计。

卸载后输出峰值约{dyn['load_release']['post_vmax']:.3f}V，电感电流最低约{dyn['load_release']['post_il_min']:.3f}A，未触发本模型11V或−1.8A保护。只代表当前控制参数与近似器件模型。

短路工况：模型故障锁存时间={1000*by['short_no_software']['fault_time']:.3f}ms，相对短路时刻约{1000*(by['short_no_software']['fault_time']-.08):.3f}ms；电感全程峰值={by['short_no_software']['il_abs_peak_all']:.3f}A。此延时包含控制响应和占空比限制等影响，不能解释为比较器本身的传播时间。输出电容短路放电脉冲达到模型值{by['short_no_software']['i_load_peak_all']:.1f}A，该值强烈依赖未给定ESR、走线和负载寄生，不能当成实物预测。保护只能停止后续能量输入，不能撤销已发生的电容放电。该次硬件行为模型触发不证明任意短路均可靠触发，也没有验证软件短路锁存、自动重启或带残压重启。

独立DC电阻网络扫描得到VREF_POS=2.55V、VREF_NEG=0.75V、VREF_OV=1.00V，比较极性对应约+1.8A/−1.8A与11V。电流扫描步长1mA、电压扫描步长1mV。HC74清零优先、ARM恢复、MCU_EN关断仅做16种理想逻辑组合核对，不是实际IC时序测试。

3V3_A由3.3V经10Ω供给，不能固定视为3.300V。举例：支路电流5mA时降压50mV，实际门限约±1.773A、10.833V，电感采样零点约1.625V。此5mA是说明性假设，需测量实际支路电流。R11/R12各1%以及输出分压各0.1%时，在3V3_A精确3.3V且不计比较器误差的条件下，过压约10.83～11.18V；电源误差还会进一步移动门限。

ADC滤波应计分压戴维南电阻：输入电压通道τ≈47.43us/fc≈3.36kHz，输出电压通道τ≈8.97us/fc≈17.74kHz，电流通道τ≈4.7us/fc≈33.86kHz。固件需要按实际参考电压、零偏和增益标定，并同步采样；不能假定五路都是4.7us。

## 6. 效率：仅能给出损耗预算

25V/8Ω：外部负载输出功率={p['pout']:.4f}W；近似功率级输入={p['pin_model']:.4f}W；该模型效率={p['eta_power_stage_model']*100:.2f}%。这里已包含模型的D1、导通电阻、分流器、DCR与ESR损耗，不含辅助电源和实际开关/磁芯损耗。

按85%要求，所有尚未建模的额外输入损耗总额最多约{budget:.3f}W。整机估算式：η=Pout/(Pin_model+Paux_input+ΔP_switch_core_thermal)。Paux_input应包含M1/M2供给驱动、栅极充放电、MCU、显示和采样的输入功耗，避免再次重复计算栅极能量。额外合计0.5W时约{100*p['pout']/(p['pin_model']+.5):.2f}%；合计1.0W时约{100*p['pout']/(p['pin_model']+1):.2f}%，后者低于85%。

作为敏感性示例，若实际开关电压电流交叠时间总和为100ns或500ns，在25V、1A、100kHz下粗估0.125W或0.625W；这只是0.5·Vin·I·(tr+tf)·f近似，不能由模型中的47ns控制RC推出。需要结合IR2104驱动能力、IRF540实际栅荷与实测波形确定。BOM没有指定更高效的功率开关或辅助模块，本次未擅自替换。

## 7. 数值检查、交付与下一步

25V/8Ω最大步长100ns→50ns，输出均值差={abs(p['vout_mean']-c['vout_mean'])*1000:.3f}mV，外部电流差={abs(p['i_load_mean']-c['i_load_mean'])*1e6:.2f}uA，说明该工况的平均值对本次步长减半不敏感；不能据此保证真实高频尖峰正确。50/150kHz敏感性下均实现约1A，但纹波与死区损耗不同，因此PWM频率必须在真实程序中明确。

附件包含报告PDF/DOCX/Markdown、可重跑Python、生成的SPICE网表、原始求解日志、JSON/CSV数值汇总与波形图。轻量压缩包不含体积较大的NPZ全波形；仓库results/data目录保留这些数据，可重跑生成。网表由simulate.py按所在路径重新生成，运行方法见README。

建议接下来按BOM选定L1/C3/C4/M1/M2具体订货型号并测量参数，明确PWM及离散控制程序，再补真实器件/自举与故障时序验证。台架按15V限流上电，逐步测试空载、20Ω、8Ω、5Ω及25V；效率必须按题意从输入台式电源端计入整机功耗，显示与按键设置另测。没有据此报告宣布现有PCB通过审查或可接真实电池。
'''
    (REPORTS/'G题_RevB_仿真验证报告.md').write_text(findings)
    render(findings)
    archive=DOWNLOADS/'G题_RevB_仿真验证资料.zip'
    with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
        for pth in ROOT.rglob('*'):
            if not pth.is_file() or pth==archive:continue
            if pth.suffix in ['.npz','.dat','.pyc','.zip'] or '__pycache__' in pth.parts:continue
            if 'pilot' in pth.name or pth.name in ['run.log','sensitivity_run.log','dynamic_rerun.log']:continue
            z.write(pth,pth.relative_to(ROOT))
    print(archive)

def render(md):
    from docx import Document
    from docx.shared import Inches,Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,Image,PageBreak
    from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.lib import colors
    doc=Document();doc.styles['Normal'].font.name='Noto Serif CJK SC';doc.styles['Normal'].font.size=Pt(10)
    doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC')
    pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
    styles=getSampleStyleSheet()
    body=ParagraphStyle('CN',fontName='STSong-Light',fontSize=10,leading=16,wordWrap='CJK',spaceAfter=7)
    heading=ParagraphStyle('CNH',parent=body,fontSize=14,leading=21,spaceBefore=10)
    story=[];lines=md.splitlines();i=0
    while i<len(lines):
        s=lines[i].strip();i+=1
        if not s:continue
        if s.startswith('|'):
            rows=[]
            while True:
                if '---' not in s:rows.append([c.strip() for c in s.strip('|').split('|')])
                if i>=len(lines) or not lines[i].strip().startswith('|'):break
                s=lines[i].strip();i+=1
            t=doc.add_table(rows=1, cols=len(rows[0]));t.style='Light Shading Accent 1'
            for j,c in enumerate(rows[0]):t.rows[0].cells[j].text=c
            for row in rows[1:]:
                cells=t.add_row().cells
                for j,c in enumerate(row):cells[j].text=c
            tbl=Table([[Paragraph(html.escape(c),body) for c in row] for row in rows],colWidths=[44,58,73,67,112,60],repeatRows=1)
            tbl.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.3,colors.grey),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#dfecf5')),('VALIGN',(0,0),(-1,-1),'TOP')]))
            story+=[tbl,Spacer(1,10)];continue
        if s.startswith('#'):
            lev=len(s)-len(s.lstrip('#'));txt=s.lstrip('#').strip();doc.add_heading(txt,level=min(lev,3));story.append(Paragraph(html.escape(txt),heading))
        else:doc.add_paragraph(s);story.append(Paragraph(html.escape(s),body))
    for name,title in [('transients.png','动态波形（假设控制环）'),('ripple.png','25V/8Ω最后100us波形'),('short.png','10mΩ短路与模型故障锁存'),('efficiency_budget.png','额外损耗预算：红线为85%边界')]:
        doc.add_heading(title,level=2);doc.add_picture(str(FIGURES/name),width=Inches(6))
        story.append(PageBreak());story.append(Paragraph(title,heading))
        from PIL import Image as PILImage
        w,h=PILImage.open(FIGURES/name).size
        story.append(Image(str(FIGURES/name),width=480,height=480*h/w))
    doc.save(REPORTS/'G题_RevB_仿真验证报告.docx')
    SimpleDocTemplate(str(REPORTS/'G题_RevB_仿真验证报告.pdf'),pagesize=(595.28,841.89),leftMargin=48,rightMargin=48,topMargin=42,bottomMargin=42).build(story)

if __name__=='__main__':main()
