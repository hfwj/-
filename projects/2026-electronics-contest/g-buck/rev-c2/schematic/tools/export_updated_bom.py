"""Produce a purchasing BOM from the delivered Rev C.2 PCB review, without edits
to the schematic, PCB, firmware, or simulation. Quantities are per one assembly.
"""
from pathlib import Path
import collections,csv,hashlib,json,re,sys,zipfile
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.worksheet.table import Table,TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.properties import CalcProperties

REV=Path(__file__).resolve().parents[2]
PCB=REV/'pcb/hardware-fit-120x90'
DATA=REV/'results/data/bom-2026-10-27'
REPORT=REV/'reports/updated-bom-2026-10-27'
DOWNLOAD=REV/'downloads'
SOURCE=PCB/'data/PCB_BOM_RevC2.csv'
DEST=DOWNLOAD/'GBUCK_RevC2_updated_BOM.xlsx'

# Explicit grouping avoids treating different precision, voltage or package
# requirements as interchangeable. Ratings labelled "建议" are procurement
# choices within the source requirements, not new edits to native design files.
GROUPS=[
('芯片','IR2184PbF','DIP-8，300mil','U1','用户已确认IR2184直插','与既有DIP8_300一致；保留设计PbF标注，不把家族确认当成完整订货号/厂家确认；库存数量另填；IR2104不可代装'),
('芯片','INA240A1D，增益20','SOIC-8','U2 U4','型号/实物核对','必须A1D；不要买成PW/TSSOP或其他增益'),
('芯片','INA181A1DBV，增益20','SOT-23-6','U3','核对已购物料','必须A1、DBV封装'),
('芯片','LM393DR（原图标LM393）','SOIC-8','U5 U9','核对已购物料','DR为参数表中的完整采购型号，开集电极输出'),
('芯片','SN74HC74DR','SOIC-14','U6','核对已购物料','不能与74HC08/74HC14混装'),
('芯片','SN74HC08DR','SOIC-14','U7','核对已购物料','与源图逻辑门型号一致'),
('芯片','SN74HC14DR','SOIC-14','U8','核对已购物料','施密特反相器'),
('功率器件','IRF540NPbF（IRF540N），N沟道100V','TO-220-3，直立','QH QL','核对已购物料','1G、2D、3S；背板是漏极；散热按实测确定'),
('功率器件','220μH±10%；Isat≥3A、Irms≥1.5A、DCR≤0.15Ω','65125立式；20×10mm投影；P10','L1','用户已确认立式；电气额定待核','沿用当前立式布局；图示6A Max不是Isat证明'),
('分流电阻','50mΩ，1%，≥1W（精度/功率仍待核）','M型康铜丝直插；图示P10、线径0.8、体高12mm','RS_IN RS_OUT','脚距/线径已获商品图支持','达科峰图标50mΩ/5A；已有备料可抵扣；连续功率、精度及平面厚度待核；独立Kelvin取样'),
('分流电阻','25mΩ，1%，≥1W（精度/功率仍待核）','U型直插；图示P5、线径0.7、体高7mm','RS_L','25mΩ商品图已给尺寸；库存待核','DILIN迪林商品图；P10占位改P5；图未给精度/功率及实际焊脚径/深度；B3已列，库存未知'),
('功率电阻','22Ω，5W，单次脉冲耐量≥0.3J','外置；主板仅P5.08引出焊盘','R_PRE','待选完整型号','器件实体外置固定；不按引出焊盘外形购买电阻'),
('电解电容','470μF/50V，低ESR，105℃','径向直插D10×H18mm，P5','C1','完整料号/纹波待核','脚径0.6mm；板上预留Dmax10.5；ESR/纹波条件待核'),
('电解电容','470μF/35V，低ESR，105℃','径向直插D10×H13mm，P5','C3','完整料号/纹波待核','核对实物耐压及极性；不要混用50V器件的高度记录'),
('贴片电容','1μF/50V，X7R','0805','C2','核对已购物料','本轮由旧C1206改0805；核对直流偏压后的有效容量'),
('贴片电容','1μF/≥25V，X7R','0805','C_BOOT C_A1 C_DRV','核对已购物料','C_A1原图仅写1μF；建议与自举/驱动统一用≥25V；不降低源图额定'),
('贴片电容','10μF/50V，X7R','0805','C4','偏压后容量待核','PCB复核版由35V提高到实购50V；必须核对有效容量'),
('贴片电容','10μF/25V，建议X7R','0805','C6 C_A3','两只0805已确认；完整料号/介质待核','共2只；C6来自原参数表，C_A3本轮由用户明确0805并替换原1206占位；实际介质与偏压后容量待核，库存另填'),
('贴片电容','100nF/50V，X7R','0805','C5 C_U2 C_U3 C_U4 C_A2 C_REF_POS C_REF_NEG C_REF_OV C_U5 C_U9 C11 C_U6 C_U7 C_U8 C_REF_IOUT C_IOUT_PROT','100nF/50V/0805/X7R，已购50只','用户对数量回复50；按单板用16只，补购0只，装一板后余34只。完整料号/容差未提供；472批次独立记4.7nF，不混计库存'),
('贴片电容','CL21B472JBANNNC；4.7nF/50V，±5%，X7R','0805','C7 C8 C9 C10','CL21B472JBANNNC，已购50只','用户对这批X7R数量回复50；0805/4.7nF/50V/±5%/X7R，单板4只，补购0只，装一板后余46只。不得用于C_VINP_ADC；与另购C0G库存2只分开'),
('贴片电容','4.7nF/50V，C0G/NP0','0805','C_VINP_ADC','4.7nF/50V/C0G/0805，已购2只','用户对C0G批次数量回复2；单板1只，补购0只，装配一板后余1只。与C7～C10的CL21B472JBANNNC/X7R库存分开；完整料号/容差未提供'),
('贴片电容','1nF/50V，C0G/NP0','0805','C_IL_ADC C_IL_PROT','1nF/50V/C0G/0805，已购4只','用户对这批0805/C0G的数量回复4；正常装配共2只，单板补购0只，装一板后余2只。与图片1206/X7R/50只分开；完整料号/容差未提供，不覆盖C_SN的100V/DNP项'),
('贴片电容','1nF/100V，C0G/NP0，DNP','0805','C_SN','DNP，默认不装','吸收电路预留，实测振铃后决定是否装配及调值'),
('普通电阻','4.7Ω，建议1%、≥0.125W','0805','RG_H RG_L','核对已购物料','不要沿用Rev B中的旧栅极电阻值'),
('普通电阻','1Ω，建议1%、≥0.125W','0805','R_VDDA','核对已购物料','模拟电源串联电阻，不能错装成0Ω'),
('普通电阻','100Ω，建议1%、≥0.125W','0805','R_PWM','核对已购物料','与100Ω/0.1%精密电阻分开采购'),
('普通电阻','1kΩ，1%、≥0.125W（统一采购建议）','0805','R3 R6 R7 R8 R_IL_ADC R_VINP_ADC R_IL_PROT R_IOUT_PROT','核对已购物料','共8只；两项保护电阻明确要求1%；原图其余未明确精度，建议统一1%'),
('普通电阻','10kΩ，建议1%、≥0.125W','0805','R_GSH R_GSL R_IN_PD R_SD_PD R12 R13 R14 R_EN_PD R_ARM_PD R_IOUT_PU','核对已购物料','普通10kΩ共10只；不替代0.1%分压/参考电阻'),
('普通电阻','23kΩ，建议1%、≥0.125W','0805','R11','核对已购物料','按23kΩ采购；不直接换成22kΩ'),
('精密电阻','100kΩ，0.1%，≥0.125W','0805','R1 R_VINP_TOP','精度/料号待核','温漂未在源图固定；按实购物料记录，不虚构已确认参数'),
('精密电阻','10kΩ，0.1%，≥0.125W','0805','R2 R4 R9 R16 R_IOUT_BOT R_VINP_BOT','精度/料号待核','共6只，与普通10kΩ区分'),
('精密电阻','1kΩ，0.1%，≥0.125W','0805','R5','精度/料号待核','与普通1kΩ区分'),
('精密电阻','49.9kΩ，0.1%，≥0.125W','0805','R10A R15A','精度/料号待核','保留49.9kΩ，不用近似值替代'),
('精密电阻','100Ω，0.1%，≥0.125W','0805','R10B R15B','精度/料号待核','共2只，与R_PWM分开'),
('精密电阻','20kΩ，0.1%，≥0.125W','0805','R_IOUT_TOP','精度/料号待核','输出告警参考上臂'),
('功率电阻','2.2kΩ，≥0.25W','1206','R_BL','功率/料号待核','确认实物额定功率，不能买成普通0805'),
('功率电阻','10Ω，≥0.25W，DNP','1206','R_SN','DNP，默认不装','与C_SN配套；由旧0805占位改1206'),
('二极管','SR560，目标60V/5A（用户选型）','SMB / DO-214AA，贴片','D1 D_FREE','用户已确认SMB；厂家资料待最终复核','共2只；1=K色条端、2=A。沿用原SMB焊盘；SR560为用户确认选型，原轴向目录不作为SMB厂家尺寸/电气手册；完整订货型号/额定条件/推荐焊盘待归档'),
('二极管','UF4007，快恢复','DO-41，成型P10.16待核','D_BOOT','成型脚距待核','工程定义1=A、2=K；不是普通1N4007'),
('二极管','BAT54S，串联双肖特基','SOT-23','D_ADC1 D_ADC2 D_ADC3 D_ADC4 D_ADC5 D_ADC6','核对已购物料','共6只；1=A1、2=K2、3=K1/A2；不能装BAT54/BAT54C'),
('二极管','BAT54，单肖特基','SOT-23','D_POR','核对已购物料','1=A、3=K、2=NC；不能与BAT54S混装'),
('保险','2A快断保险管，5×20mm，商品标250V','配5×20mm管夹','F1','尺寸/快断已给资料；实际电气性能待核','商品BXG16519；250V未注明AC/DC分断条件，启动不误熔断与故障分断待验证；管夹另列2只'),
('连接器','2P接线端子，P2.54，≥3A','图示5.48×6.30×8.80mm；孔1.30','J1 J2 J_BYP J_MON','完整型号/朝向待核','共4只；照片标称6A/150V；J_BYP是接口，不包含外接开关'),
('连接器','1×4直排针，P2.54','直插，针长/塑高待核','J_AUX12 J_AUX5','配对件待核','共2只；与LM2596HV通过排线相接，不直接按四角焊盘插装'),
('连接器','2×8排针，P2.54','直插；配对高度待核','J_CTL','配对件待核','1只，16针；PA6监测走独立J_MON，不占此16针'),
('测试点','裸铜测试焊盘，直径1.5mm','PCB焊盘；默认不买测试柱','TP1 TP2 TP3 TP4 TP5 TP6 TP7 TP8 TP9 TP10 TP11 TP12 TP13 TP14','默认无需采购','14个设计测试点，实物采购数量为0；若改测试柱需同步封装'),
]

EXTERNAL=[
 ['外接模块','M1','LM2596HV降压模块，输出调至12V，负载能力≥0.2A',1,'块','正文标称输入5～55V，覆盖15～25V；外接J_AUX12；先调压后接负载'],
 ['外接模块','M2','LM2596HV降压模块，输出调至5V，负载能力≥0.5A',1,'块','与M1是两块独立模块；外接J_AUX5；额定范围不等于带载实测'],
 ['外接模块','MCU','STM32F103C8T6核心板',1,'块','继续排线外接；确认5V输入、3.3V输出及实际排针；芯片手册不代表核心板机械图'],
 ['外接模块','LCD','1.8英寸ST7735类SPI显示模块',1,'块','附件是11Pin版本，核对实际版本及供电跳线；SCK=PA5、MOSI=PA7；PA6留给ADC'],
 ['保险配件','F1管夹','5×20mm黄铜镀镍管夹；单夹脚中心距5.66mm',2,'只','用户确认5.66是中心距；脚宽1.1/厚0.5；槽0.9×1.5设计值；双夹中心距15仅试装初值，旧6×30尺寸作废'],
 ['外接开关','J_BYP配套','常开保持型SPST旁路开关，≥30VDC/≥3A',1,'只','与PCB端子分开计数；需明确直流额定，不能仅凭交流额定替代'],
 ['配线','J_AUX12/J_AUX5','P2.54，1×4配套母座带线',2,'套','按VIN、GND、OUT、GND四线接；线长由模块安装位置确定'],
 ['配线','J_CTL配套','P2.54，2×8配套母座及16芯控制线',1,'套','按逐引脚表定制；核对1脚与插接方向，不假设核心板可直接插装'],
 ['配线','J_MON配套','PA6/GND两线连接线',1,'套','接J_MON.1→PA6、J_MON.2→GND；与LCD PA5分开'],
 ['安装件','PCB固定','铜柱、螺丝、垫圈',None,'套','数量/高度待外壳核对；当前4个拟定安装孔不代表安装方案已经确认'],
 ['安装件','QH/QL散热','散热片及必要的绝缘导热垫/肩套',None,'套','按温升与固定方式选配；金属背板是漏极，不默认共用导电散热片'],
 ['配线','R_PRE/电源线','外置功率电阻固定件及电源导线',None,'套','R_PRE电阻已在主清单计1只，不再重复购买；长度/线径按布局与电流确认'],
]

# User-confirmed quantity; other groups remain unknown, not zero.
CONFIRMED_STOCK = {'100nF/50V，X7R': 50, '4.7nF/50V，C0G/NP0': 2, 'CL21B472JBANNNC；4.7nF/50V，±5%，X7R': 50, '1nF/50V，C0G/NP0': 4}

NOTES=[
 ['最新1nF数量确认', '比较6f26307：用户对另购1nF/50V/C0G（NP0）/0805数量回复“4”，登记库存4只。C_IL_ADC、C_IL_PROT单板各1只，共2只；需补购0只，预计装一板后余2只，尚未实际扣减库存。Excel库存填4、补购保留动态公式，CSV写当前4/0。图片中的1206/X7R/50只另记，不混入本项；本次50V备料不覆盖C_SN的100V/C0G/DNP预留。完整料号/容差待归档；采购46行、120逻辑位号/104默认有料及先前三类库存保持。原理图/PCB原生文件、网络、脚号、设计值/封装及单板用量未改，仅更新库存和补购记录，无新增器件或改板要求；仿真/固件未修改或运行。'],
 ['历史1nF另购确认（数量已后补）', '比较1dac34a：用户对是否另有1nF/50V/C0G（NP0）/0805回复“买了”，确认正常装配C_IL_ADC、C_IL_PROT所需C0G款已另购，每板共2只。实际数量、完整料号和容差未提供，库存/补购栏留空，不把设计需求2只当成库存。此前图片1206/1nF/±10%/50V/X7R/50只仍单独登记，不抵扣0805/C0G需求；本轮50V规格确认不用于C_SN的1nF/100V/C0G/DNP预留。原理图和PCB原生文件、网络、脚号、设计值/封装及用量均未改；仅更新BOM已购状态，不需要本轮改板或新增逻辑器件。46采购行、120逻辑位号、104默认有料及先前三类已确认库存保持；仿真/固件未修改或运行。'],
 ['历史1nF图片核对（后续另购见最新项）', '比较8d57b27：用户新图已选规格为1206、1nF、±10%、50V、X7R，50只；无完整厂家型号。原理图C_IL_ADC/C_IL_PROT正常装配共2只，PCB均C0805，且C_IL_PROT明确C0G。这批1206/X7R不能视作0805/C0G可直接装配库存，50只不填入目标行库存。建议另备1nF/50V/C0G/0805共2只以统一备料；是否已有该款待核，不自动认定缺货2只。C_SN另为1nF/100V/C0G、DNP，本图也不能用于确认该项。原理图/PCB原生文件、网络、设计值和封装均未改；没有放大焊盘或放宽介质要求。仅更新实购物料记录及BOM待核状态，仿真/固件未修改或运行。'],
 ['版本基准', '本次比较6f26307；1nF/50V/C0G/0805已购4只，单板用2只/补购0只/装一板余2只；仅更新库存，设计保持。'],
 ['计数口径','120个逻辑位号含14个裸测试点、2个DNP预留和1个外置R_PRE。120不等于需购买120个独立实体。'],
 ['默认装配','120−14裸测试点−2个DNP=104个有料逻辑位号；含R_PRE外置电阻。管夹、模块、开关和线束另表计数。'],
 ['库存填写', 'G列已填用户确认库存：100nF/X7R 50只、4.7nF/X7R 50只、4.7nF/C0G 2只、1nF/C0G 4只，四项单板补购均0。其他空白表示未知，不代表未购。'],
 ['DNP','R_SN/C_SN各设计1个，默认装配0个；实测振铃后决定是否装配和调值。'],
 ['采购规格','标“建议/统一采购”的额定、精度或介质是采购建议，区别于源图原始标注；实际原始值在逐位号页保留。'],
 ['普通/精密电阻功率','除源图另有功率要求的R_BL/R_SN/R_PRE和分流器外，0805电阻统一≥0.125W是本表采购建议，不冒充源图已规定的功率。'],
 ['优先核对','U1型号家族/直插类别已确认；保留IR2184PbF设计标注，不再列为型号封装待确认或默认补购；库存数量另填，IR2104不可代装。RS_L 25mΩ在B3清单中已包含，旧表记未购，现库存待核，不再自动假定库存为0。'],
 ['电容数量','100nF共16只、4.7nF共5只、默认1nF共2只，另有1只DNP的1nF/100V。'],
 ['料号纠错','472批次已确认4.7nF；随后用户确认另外买了100nF/0805，两批分开记录。100nF/50V/0805/X7R已购50只，单板需求16只，补购0只，装配一板后余34只；完整料号/容差未提供；用户现已确认实际型号CL21B472JBANNNC，先前口头C0G由实际型号纠正为X7R；该款50V/±5%/0805按用户编码图明确。C_VINP_ADC所需的4.7nF/50V/C0G/0805另购状态已获确认；该批C0G库存2只、单板用1只/补购0只；C0G完整料号/容差待归档；CL21B472JBANNNC/X7R库存现已确认50只、单板用4只/补购0只，两类不混计。'],
 ['分流器新证据','RS_IN/OUT：50mΩ P10、图示线径0.8、体高12、标称5A；RS_L：25mΩ P5、图示线径0.7、体高7。电阻体线径不自动证明焊接直脚径，精度/连续功率仍待核。'],
 ['保险夹确认','用户明确单夹脚中心距5.66mm；脚宽1.1/材料厚0.5。槽0.9×1.5及双夹中心距15mm是设计选择，装入5×20实物核对端帽接触/止挡后方可冻结；不能使用旧6×30/7.7尺寸。'],
 ['二极管确认','两只SR560/SMB已由用户明确；沿用现有SMB标准几何，1=K/2=A。DO-201AD图不用于贴片建库，厂家完整电气手册/推荐焊盘最终复核仍需归档。'],
 ['未确认封装','C_A3的0805类别已确认；F1组合试装、分流器完整包络及外壳固定仍待核。C_A3完整料号/介质/偏压有效容量、二极管厂家数据待最终复核。'],
 ['电感','65125图示6A Max不是饱和电流保证；源设计仍要求Isat≥3A、Irms≥1.5A、DCR≤0.15Ω。'],
 ['当前制造状态','PCB仍有未布连接，客户端打开、原生ERC/DRC未完成；清单提供采购核对依据，不是制板放行。'],
 ['本次原理图影响', '相对6f26307，原理图和PCB原生文件未改，网络/脚号/设计值/封装与单板用量保持。'],
 ['本次器件/PCB影响', '1nF/50V/C0G/0805登记库存4只，单板用2只/补购0只，预计装一板余2只；1206/X7R款分开。完整料号/容差待归档，46采购行/120位号/104默认有料及既有三类库存保持，无新增或改板要求。'],
 ['本次仿真/固件影响','未修改或运行Multisim/ngspice/STM32固件，旧器件/功能等效模型未同步为SR560厂家模型，不证明新选型压降、温升、效率或保护实测通过。'],
]

def writecsv(path,header,rows):
    with path.open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f,lineterminator='\n');w.writerow(header);w.writerows(rows)

def main():
    for p in [DATA,REPORT,DOWNLOAD]:p.mkdir(parents=True,exist_ok=True)
    with SOURCE.open(encoding='utf-8-sig') as f:source=list(csv.DictReader(f))
    original={r['位号']:r for r in source}
    assert len(original)==len(source)==120
    groups=[];seen=[];detail=[]
    for cat,spec,package,refs,status,note in GROUPS:
        rr=refs.split();seen.extend(rr);quantity=len(rr)
        default=0 if cat=='测试点' or 'DNP' in spec else quantity
        stock=CONFIRMED_STOCK.get(spec)
        needed=max(default-stock,0) if stock is not None else None
        groups.append([cat,spec,package,'、'.join(rr),quantity,default,stock,needed,status,note])
        for ref in rr:
            r=original[ref]
            itemnote=note
            detail.append([ref,r['数值/型号'],r['PCB封装'],spec,package,1,0 if default==0 else 1,
                           status,itemnote,r['装配与核查说明']])
    assert len(seen)==len(set(seen))==120
    assert set(seen)==set(original)
    assert sum(r[4] for r in groups)==120
    assert sum(r[5] for r in groups)==104
    counts=collections.Counter()
    for row in source:
        if not row['位号'].startswith('C'):continue
        value=row['数值/型号']
        match=re.match(r'(100nF|4\.7nF|1nF)',value)
        if match:counts[(match[1],'DNP' in value)]+=1
    assert counts=={('100nF',False):16,('4.7nF',False):5,('1nF',False):2,('1nF',True):1}
    # Independently compare all source values/packages to saved native schematic.
    from native_common import read_project,inventory,sha
    native=PCB/'native/GBUCK_RevC2_120x90_PCB_review.eprj2'
    m,d,_=read_project(native)
    _,_,parts,pins=inventory(m,d,'G题同步Buck Rev C.2 原理图')
    assert set(parts)==set(original)
    for ref,r in original.items():
        assert parts[ref]['attrs']['Value']==r['数值/型号'],ref
        fp=parts[ref]['attrs']['Footprint']
        assert d[fp]['META','META']['title']==r['PCB封装'],ref
    wb=Workbook();wb.remove(wb.active);wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
    ws=wb.create_sheet('使用说明');ws.append(['项目','说明'])
    for row in NOTES:ws.append(row)
    ws=wb.create_sheet('采购汇总')
    gh=['类别','采购规格/型号','采购封装/尺寸','对应位号','设计数量','默认装配数量','现有可用数量_填写','需补购数量_自动','资料状态','采购与封装说明']
    ws.append(gh)
    for n,row in enumerate(groups,2):
        ws.append(row);ws.cell(n,8,f'=IF(G{n}="","",MAX(F{n}-G{n},0))')
        ws.cell(n,7).fill=PatternFill('solid',fgColor='FFF2CC')
    dv=DataValidation(type='whole',operator='greaterThanOrEqual',formula1=0,allow_blank=True)
    dv.errorTitle='请输入非负整数';dv.error='库存数量应为0或正整数';dv.showErrorMessage=True
    ws.add_data_validation(dv);dv.add(f'G2:G{len(groups)+1}')
    ws=wb.create_sheet('逐位号_120项')
    dh=['位号','当前工程原始数值/型号','当前PCB封装','采购规格/建议','采购封装说明','设计数量','默认装配数量','资料状态','采购说明','原PCB_BOM备注']
    ws.append(dh)
    for row in detail:ws.append(row)
    ws=wb.create_sheet('外接模块与辅料')
    eh=['类别','配套对象','规格','数量','单位','说明']
    ws.append(eh)
    for row in EXTERNAL:ws.append(row)
    ws=wb.create_sheet('重点变化与待确认')
    ws.append(['对象','旧表/原始资料','当前PCB采购基准','必须处理'])
    changes=[
      ['U1','Rev B表：IR2104PbF','设计IR2184PbF，DIP8','用户已确认IR2184直插；无需本轮改板，不默认补购；完整订货号/库存数量另记'],
      ['RS_L','旧P10占位；库存未知','25mΩ U型；P5、图示线径0.7、体高7','PCB改P5；精度、连续功率及实际焊脚径/深度仍待核'],
      ['RS_IN/RS_OUT','P10原为占位；原设计R2512','50mΩ M型，商品图P10/线径0.8/体高12/额定5A','脚距已确认；精度、连续功率和包络厚度待核'],
      ['F1','中间采用6×30/7.7，现作废','5×20mm、2A快断；配2只P5.66夹','脚宽1.1/厚0.5；四槽更新；双夹15mm试装初值必须实物核对'],
      ['J1/J2/J_BYP/J_MON','旧P5.08端子','P2.54、2P端子，共4只','核对具体型号和进线方向'],
      ['C1/C3','旧D12.5占位','D10、P5；高18/13分别记录','完整料号、纹波和实际耐压待核'],
      ['C2/C4/C6/C_A3','原C1206占位','现均0805；C4为10μF/50V','C_A3本轮由用户单独确认0805，10μF/25V保持'],
      ['100nF电容','原472批次实际4.7nF；另购100nF/50V/0805/X7R共50只','Rev C2单板需100nF共16只','本项补购0只，装配一板后余34只；完整料号/容差未提供'],
      ['4.7nF电容','C7～C10实购CL21B472JBANNNC/X7R；现另购C0G款也获确认','C7/C8/C9/C10用X7R共4只；C_VINP_ADC用C0G 1只','X7R库存50只，单板4只/补购0只；C0G库存2只，单板1只/补购0只，两类不混装；C0G完整料号/容差待归档'],
      ['R_SN/C_SN','吸收预留','各1，默认不装；R_SN为1206','DNP不列为默认必装'],
      ['测试点','TP1～TP10旧表','TP1～TP14共14个裸焊盘','默认无需采购测试柱'],
      ['J_MON','C.1 PA5','C.2 PA6/ADC12_IN6','PA5保留LCD SCK'],
    ]
    for row in changes:ws.append(row)
    widths={
      '使用说明':[23,110],
      '采购汇总':[15,51,40,70,12,15,19,19,25,100],
      '逐位号_120项':[20,34,34,50,40,12,16,25,90,90],
      '外接模块与辅料':[17,26,68,12,10,110],
      '重点变化与待确认':[28,48,65,85],
    }
    for n,ws in enumerate(wb,1):
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        ws.sheet_view.zoomScale=80
        for cell in ws[1]:
            cell.font=Font(name='Microsoft YaHei',color='FFFFFF',bold=True,size=10)
            cell.fill=PatternFill('solid',fgColor='24584A')
            cell.alignment=Alignment(vertical='center',wrap_text=True)
        ws.row_dimensions[1].height=30
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font=Font(name='Microsoft YaHei',size=10)
                cell.alignment=Alignment(vertical='top',wrap_text=True)
            ws.row_dimensions[row[0].row].height=58 if ws.title not in ['使用说明','外接模块与辅料'] else 46
        from openpyxl.utils import get_column_letter
        for c,w in enumerate(widths[ws.title],1):ws.column_dimensions[get_column_letter(c)].width=w
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
        ws.print_title_rows='1:1';ws.print_options.horizontalCentered=True
    wb.save(DEST)
    writecsv(DATA/'purchase_grouped.csv',gh,groups)
    writecsv(DATA/'bom_by_reference_120.csv',dh,detail)
    writecsv(DATA/'external_modules_and_accessories.csv',eh,EXTERNAL)
    text=['# Rev C2更新物料清单','',
      '基准：当前PCB复核工程；每套用量。完整Excel含采购汇总、120项逐位号、外接辅料与待确认项。',
      '', '| 类别 | 采购规格 | 封装 | 设计数 | 默认装配数 | 位号 |', '|---|---|---|---:|---:|---|']
    for cat,spec,package,refs,q,default,_,__,status,note in groups:
        text.append(f'| {cat} | {spec} | {package} | {q} | {default} | {refs} |')
    text.extend(['','## 外接模块和辅料','','| 对象 | 规格 | 数量 | 说明 |','|---|---|---|---|'])
    for cat,target,spec,q,unit,note in EXTERNAL:text.append(f'| {target} | {spec} | {str(q)+unit if q is not None else "待定"} | {note} |')
    text.extend(['','## 使用与修改影响',''])
    for title,note in NOTES:text.append(f'- **{title}：**{note}')
    (REPORT/'物料清单.md').write_text('\n'.join(text)+'\n')
    # Readable purchasing PDF; inventory-dependent supplement quantities remain
    # in Excel and are not invented here.
    import html
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table as PDFTable,TableStyle,PageBreak
    pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
    body=ParagraphStyle('body',fontName='STSong-Light',fontSize=8.2,leading=11,wordWrap='CJK',spaceAfter=7)
    title=ParagraphStyle('title',parent=body,fontSize=16,leading=20,spaceAfter=10)
    cellstyle=ParagraphStyle('cell',parent=body,fontSize=7.5,leading=10)
    def p(t):return Paragraph(html.escape(str(t)),cellstyle)
    def table(rows,widths):
        t=PDFTable([[p(c) for c in row] for row in rows],colWidths=widths,repeatRows=1)
        t.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.3,colors.HexColor('#b8c6bf')),
            ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e7f0eb')),('VALIGN',(0,0),(-1,-1),'TOP'),
            ('TOPPADDING',(0,0),(-1,-1),3),('BOTTOMPADDING',(0,0),(-1,-1),3)]))
        return t
    flow=[Paragraph('G题 Rev C2 更新物料清单',title),
        Paragraph('每套用量；基准为当前PCB复核版。120个逻辑位号，默认有料104项，另有14个裸测试点和2个DNP。库存填写和补购计算见Excel；本表数量不是需新购买数量。',body),
        Paragraph('商品图尺寸：RS_IN/OUT 50mΩ脚距10mm、线径0.8mm；RS_L 25mΩ脚距5mm、线径0.7mm。精度/功率等仍待核。U1=IR2184PbF；100nF共16只。',body)]
    rows=[['类别','规格/型号','封装/尺寸','设计数','默认数','位号']]
    rows.extend([[r[0],r[1],r[2],r[4],r[5],r[3]] for r in groups])
    flow.extend([table(rows,[53,210,154,32,32,281]),Spacer(1,12),Paragraph('外接模块及配套辅料',title)])
    rows=[['对象','规格','数量','说明']]
    rows.extend([[r[1],r[2],str(r[3])+r[4] if r[3] is not None else '待定',r[5]] for r in EXTERNAL])
    flow.extend([table(rows,[85,250,40,387]),Spacer(1,12),Paragraph('采购与修改说明',title)])
    for label,note in NOTES:flow.append(Paragraph(html.escape(label+'：'+note),body))
    def footer(canvas,doc):
        canvas.setFont('STSong-Light',8);canvas.drawString(40,20,'Rev C2 PCB复核版采购核对清单 / 2026-10-27')
        canvas.drawRightString(800,20,str(doc.page))
    pdf=DOWNLOAD/'GBUCK_RevC2_updated_BOM.pdf'
    SimpleDocTemplate(str(pdf),pagesize=(841.89,595.276),leftMargin=40,rightMargin=40,topMargin=32,bottomMargin=35).build(flow,onFirstPage=footer,onLaterPages=footer)
    evidence={'status':'PASS_BOM_COVERAGE_AND_NATIVE_READBACK','source_csv_sha256':sha(SOURCE),
      'source_native_sha256':sha(native),'unique_logical_references':len(seen),'grouped_lines':len(groups),
      'design_quantity':120,'default_populated_logical_references':104,'bare_testpads':14,
      'DNP_references':['R_SN','C_SN'],'DNP_quantity':2,
      'cap_counts':{'100nF_default':16,'4.7nF_default':5,'1nF_default':2,'1nF_DNP':1},
      'external_accessory_lines':len(EXTERNAL),'source_files_modified':False,
      '4p7nF_selection':{'CL21B472JBANNNC_X7R':{'refs':['C7','C8','C9','C10'],'quantity':4,'stock':50},'C0G_separate_requirement':{'refs':['C_VINP_ADC'],'quantity':1,'stock':2,'full_part_number':None,'separate_purchase_confirmed':True}},
      'confirmed_inventory':{'100nF_50V_0805_X7R':{'available':50,'per_board':16,'purchase_needed_one_board':0,'remaining_after_one_board':34},'4p7nF_50V_0805_C0G':{'available':2,'per_board':1,'purchase_needed_one_board':0,'remaining_after_one_board':1},'4p7nF_50V_0805_X7R_CL21B472JBANNNC':{'available':50,'per_board':4,'purchase_needed_one_board':0,'remaining_after_one_board':46},'1nF_50V_0805_C0G':{'available':4,'per_board':2,'purchase_needed_one_board':0,'remaining_after_one_board':2}},
      'source_graph_connected_pins':sum(bool(p['net']) for p in pins),
      'manufacturing_release':False,'xlsx_sha256':sha(DEST),'pdf_sha256':sha(pdf)}
    # Read the saved workbook back: formulas, reference coverage and quantities.
    got=load_workbook(DEST,data_only=False)
    assert got['逐位号_120项'].max_row==121
    assert {row[0] for row in got['逐位号_120项'].iter_rows(min_row=2,values_only=True)}==set(original)
    assert sum(row[5] for row in got['采购汇总'].iter_rows(min_row=2,values_only=True))==104
    for row in got['采购汇总'].iter_rows(min_row=2,values_only=True):
        assert str(row[7]).startswith('=IF(')
        assert row[6]==CONFIRMED_STOCK.get(row[1])
    with (DATA/'purchase_grouped.csv').open(encoding='utf-8-sig') as f:
        saved=list(csv.DictReader(f))
    cap=next(r for r in saved if r['采购规格/型号']=='100nF/50V，X7R')
    assert (cap['默认装配数量'],cap['现有可用数量_填写'],cap['需补购数量_自动'])==('16','50','0')
    cog=next(r for r in saved if r['采购规格/型号']=='4.7nF/50V，C0G/NP0')
    assert (cog['默认装配数量'],cog['现有可用数量_填写'],cog['需补购数量_自动'])==('1','2','0')
    x7r=next(r for r in saved if r['采购规格/型号']=='CL21B472JBANNNC；4.7nF/50V，±5%，X7R')
    assert (x7r['默认装配数量'],x7r['现有可用数量_填写'],x7r['需补购数量_自动'])==('4','50','0')
    one=next(r for r in saved if r['采购规格/型号']=='1nF/50V，C0G/NP0')
    assert (one['默认装配数量'],one['现有可用数量_填写'],one['需补购数量_自动'])==('2','4','0')
    assert all(r['现有可用数量_填写']=='' and r['需补购数量_自动']=='' for r in saved if r is not cap and r is not cog and r is not x7r and r is not one)
    assert sha(SOURCE)==evidence['source_csv_sha256'] and sha(native)==evidence['source_native_sha256']
    (DATA/'verification.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2))
    print(json.dumps(evidence,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
